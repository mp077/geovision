"""
GeoVision AI — Dataset Excel Exporter

Generates rich, beautifully formatted Excel spreadsheets (.xlsx) for geological datasets.
Includes:
  - Summary sheet with dataset KPIs
  - All Images sheet (resolution, file size, hashes, labels)
  - Train / Val / Test split sheet
  - Class distribution breakdown with percentages
  - Rejected / duplicate images audit sheet
"""

import os
from datetime import datetime
from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def format_excel_sheet(ws, header_color="2B5742", accent_color="F4F2EB"):
    """Applies modern professional styling to an openpyxl worksheet."""
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
    regular_font = Font(name="Segoe UI", size=10)
    border_thin = Side(border_style="thin", color="D1D5DB")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    
    # Freeze header row
    ws.freeze_panes = "A2"
    
    # Enable auto-filter
    if ws.max_column > 0 and ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        # Format header
        header_cell = ws.cell(row=1, column=col_idx)
        header_cell.font = header_font
        header_cell.fill = header_fill
        header_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        # Calculate column width
        max_len = len(str(header_cell.value or ""))
        for row_idx in range(2, min(ws.max_row + 1, 300)):  # sample up to 300 rows for width
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = regular_font
            cell.border = cell_border
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
                
        # Set column width with safety bounds
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 65)

    # Style remaining data cells border
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.font = regular_font
            cell.border = cell_border


def export_dataset_to_excel(
    manifest_dir: Path | str = "data/manifests",
    output_excel_path: Path | str = "data/manifests/dataset_manifest.xlsx",
    all_records: list[dict] | None = None,
    split_records: list[dict] | None = None,
    class_dist: dict[str, int] | None = None,
    rejected_records: list[dict] | None = None,
) -> Path:
    """
    Exports geological dataset records to a comprehensive multi-sheet Excel file (.xlsx).
    If records are not provided in memory, reads from manifest CSV files.
    """
    manifest_dir = Path(manifest_dir)
    output_excel_path = Path(output_excel_path)
    output_excel_path.parent.mkdir(parents=True, exist_ok=True)

    # 1. Load or prepare All Images
    if all_records is not None:
        df_all = pd.DataFrame(all_records)
    elif (manifest_dir / "all_images.csv").exists():
        df_all = pd.read_csv(manifest_dir / "all_images.csv")
    else:
        df_all = pd.DataFrame()

    if not df_all.empty and "size_bytes" in df_all.columns:
        df_all["size_kb"] = (df_all["size_bytes"] / 1024).round(2)
        # Rename columns for professional display
        col_rename = {
            "path": "File Path",
            "label": "Geological Class",
            "sha256": "SHA-256 Checksum",
            "phash": "Perceptual Hash (pHash)",
            "width": "Width (px)",
            "height": "Height (px)",
            "format": "Format",
            "size_bytes": "File Size (Bytes)",
            "size_kb": "File Size (KB)"
        }
        df_all = df_all.rename(columns={k: v for k, v in col_rename.items() if k in df_all.columns})

    # 2. Load or prepare Split Manifest
    if split_records is not None:
        df_splits = pd.DataFrame(split_records)
    elif (manifest_dir / "split_manifest.csv").exists():
        df_splits = pd.read_csv(manifest_dir / "split_manifest.csv")
    else:
        df_splits = pd.DataFrame()

    if not df_splits.empty:
        split_rename = {
            "split": "Split (Subset)",
            "label": "Geological Class",
            "path": "File Path"
        }
        df_splits = df_splits.rename(columns={k: v for k, v in split_rename.items() if k in df_splits.columns})
        if "Split (Subset)" in df_splits.columns:
            df_splits["Split (Subset)"] = df_splits["Split (Subset)"].str.upper()

    # 3. Load or prepare Class Distribution
    if class_dist is not None:
        dist_items = [{"class": k, "count": v} for k, v in sorted(class_dist.items(), key=lambda x: -x[1])]
        df_dist = pd.DataFrame(dist_items)
    elif (manifest_dir / "class_distribution.csv").exists():
        df_dist = pd.read_csv(manifest_dir / "class_distribution.csv")
    else:
        df_dist = pd.DataFrame()

    if not df_dist.empty and "count" in df_dist.columns:
        total_imgs = df_dist["count"].sum()
        df_dist["percentage"] = (df_dist["count"] / total_imgs * 100).round(2).astype(str) + " %"
        dist_rename = {
            "class": "Rock / Mineral Class",
            "count": "Total Images",
            "percentage": "Share of Dataset (%)"
        }
        df_dist = df_dist.rename(columns={k: v for k, v in dist_rename.items() if k in df_dist.columns})

    # 4. Load or prepare Rejected Records
    if rejected_records is not None:
        df_rejected = pd.DataFrame(rejected_records)
    elif (manifest_dir / "rejected_images.csv").exists():
        df_rejected = pd.read_csv(manifest_dir / "rejected_images.csv")
    else:
        df_rejected = pd.DataFrame()

    if not df_rejected.empty:
        rej_rename = {
            "path": "File Path",
            "reason": "Rejection Reason",
            "duplicate_of": "Duplicate Of (Reference File)"
        }
        df_rejected = df_rejected.rename(columns={k: v for k, v in rej_rename.items() if k in df_rejected.columns})

    # 5. Create Dataset Summary Sheet
    total_valid = len(df_all)
    total_rejected = len(df_rejected)
    total_scanned = total_valid + total_rejected
    classes_count = len(df_dist) if not df_dist.empty else 0
    
    train_count = len(df_splits[df_splits["Split (Subset)"] == "TRAIN"]) if not df_splits.empty and "Split (Subset)" in df_splits.columns else 0
    val_count = len(df_splits[df_splits["Split (Subset)"] == "VAL"]) if not df_splits.empty and "Split (Subset)" in df_splits.columns else 0
    test_count = len(df_splits[df_splits["Split (Subset)"] == "TEST"]) if not df_splits.empty and "Split (Subset)" in df_splits.columns else 0

    summary_data = [
        {"Property": "Dataset Name", "Value": "GeoVision AI Geological Rock & Mineral Dataset"},
        {"Property": "Dataset Version", "Value": "1.0.0"},
        {"Property": "Last Exported", "Value": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
        {"Property": "Total Scanned Candidate Images", "Value": total_scanned},
        {"Property": "Valid Clean Images Accepted", "Value": total_valid},
        {"Property": "Rejected / Corrupt / Duplicate Images", "Value": total_rejected},
        {"Property": "Total Geological Classes", "Value": classes_count},
        {"Property": "Training Split (70%)", "Value": f"{train_count} images"},
        {"Property": "Validation Split (15%)", "Value": f"{val_count} images"},
        {"Property": "Test Split (15%)", "Value": f"{test_count} images"},
        {"Property": "Primary Classes", "Value": ", ".join(df_dist["Rock / Mineral Class"].tolist()) if not df_dist.empty else "N/A"},
        {"Property": "Format", "Value": "Microsoft Excel OpenXML (.xlsx)"},
    ]
    df_summary = pd.DataFrame(summary_data)

    # Write all dataframes into the Excel workbook
    with pd.ExcelWriter(output_excel_path, engine="openpyxl") as writer:
        df_summary.to_excel(writer, sheet_name="Dataset Summary", index=False)
        if not df_dist.empty:
            df_dist.to_excel(writer, sheet_name="Class Distribution", index=False)
        if not df_splits.empty:
            df_splits.to_excel(writer, sheet_name="Split Manifest", index=False)
        if not df_all.empty:
            df_all.to_excel(writer, sheet_name="All Images", index=False)
        if not df_rejected.empty:
            df_rejected.to_excel(writer, sheet_name="Rejected Images", index=False)

    # Re-open with openpyxl to apply rich styling
    from openpyxl import load_workbook
    wb = load_workbook(output_excel_path)
    
    # Style summary sheet
    if "Dataset Summary" in wb.sheetnames:
        ws_sum = wb["Dataset Summary"]
        format_excel_sheet(ws_sum, header_color="1E293B")
        # Bold property names
        for row in range(2, ws_sum.max_row + 1):
            ws_sum.cell(row=row, column=1).font = Font(name="Segoe UI", size=10, bold=True)
            ws_sum.cell(row=row, column=2).alignment = Alignment(horizontal="left")

    # Style other sheets
    sheet_colors = {
        "Class Distribution": "2B5742",
        "Split Manifest": "1E40AF",
        "All Images": "334155",
        "Rejected Images": "991B1B"
    }

    for name, color in sheet_colors.items():
        if name in wb.sheetnames:
            format_excel_sheet(wb[name], header_color=color)

    wb.save(output_excel_path)
    print(f"[OK] Master dataset Excel file generated successfully: {output_excel_path}")

    # Also create convenient copies at data/dataset.xlsx and data/manifests/dataset.xlsx
    copy_paths = [
        Path("data/dataset.xlsx"),
        Path("data/manifests/dataset.xlsx"),
        Path("data/manifests/all_images.xlsx")
    ]
    for cp in copy_paths:
        try:
            cp.parent.mkdir(parents=True, exist_ok=True)
            import shutil
            shutil.copy2(output_excel_path, cp)
            print(f"     Copy created at: {cp}")
        except Exception as e:
            print(f"     Could not copy to {cp}: {e}")

    return output_excel_path


if __name__ == "__main__":
    export_dataset_to_excel()
